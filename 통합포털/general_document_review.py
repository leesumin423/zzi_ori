# -*- coding: utf-8 -*-
"""전자결재(그룹웨어 기안) 문서를 결재 상신 "전에" 미리 훑어보는 범용 AI 검수.

disclosure_review/ai_review.py는 (주)동양 정기공시 전용(작성지침서 기준 충족
여부)이라 다른 부서ㆍ다른 종류의 기안 문서에는 못 쓴다 — 이 모듈은 그 대신
어떤 기안 문서든(회사 지정 작성기준이 없어도) 본문 텍스트 + 첨부파일을 함께
받아 아래를 봐준다(EGCHECK가 보는 4가지 중 ①②③, ④연결문서 정합성은 아직
없음 — 그룹웨어에서 실제로 문서를 읽어와야 해서 별도로 다뤄야 한다):
  ① 본문 안에서 서술문ㆍ수치가 서로 모순되지 않는지, 오탈자ㆍ산술 오류
  ② 첨부파일(docx/xlsx/pdf) 안에서도 마찬가지로 자기모순이 없는지
  ③ 본문에서 말한 내용(금액ㆍ수량ㆍ날짜 등)이 첨부파일 내용과 실제로 맞는지

호출 방식은 disclosure_review/ai_review.py와 동일하게 이 PC에 설치된 Claude
Code CLI를 헤드리스로 부른다(구독 한도 안에서 동작, 사람이 버튼을 눌렀을
때만 1회 실행) — 두 모듈이 거의 같은 서브프로세스 호출 로직을 각자 갖고
있는데, 이건 disclosure_review 패키지 안의 검증된 코드를 건드리지 않고
독립적으로 유지하기 위한 의도적인 중복이다(공시 전용 검수가 잘못될 리스크를
이 기능이 절대 만들지 않도록).
"""
import json
import os
import shutil
import subprocess
import tempfile

# disclosure_review/ai_review.py와 같은 이유로 시간 제한을 두지 않는다 —
# 목적이 "빨리"가 아니라 "정확하게"라서, 일찍 끊어 재시도하게 하는 것보다
# 끝까지 기다리는 쪽이 사용량ㆍ정확도 모두에 낫다.
CLI_TIMEOUT_SEC = None
# Read는 허용한다 — 스캔본(이미지형) PDF처럼 텍스트 추출이 안 되는 첨부파일은
# 파일을 격리 폴더에 그대로 저장해두고 Claude가 자기 Read 도구(비전 포함)로
# 직접 읽게 한다. 그 외(Bash/Write/Edit 등 수정·실행 계열)는 여전히 다 막는다.
_CLI_DISALLOWED_TOOLS = 'Bash Write Edit Glob Grep WebFetch WebSearch NotebookEdit Agent Artifact'
_CLAUDE_CLI_PATH = ''  # '' = 아직 탐색 안 함

# 첨부파일 1개당 넣을 텍스트 상한 — 큰 엑셀ㆍPDF 하나가 프롬프트 전체를
# 잡아먹어 다른 첨부ㆍ본문 검토를 방해하지 않도록 자른다. 잘린 경우 표시해서
# "이 파일은 일부만 봤다"는 걸 결과에서 알 수 있게 한다.
MAX_CHARS_PER_ATTACHMENT = 30000
SUPPORTED_ATTACHMENT_EXTS = ('.docx', '.xlsx', '.xlsm', '.pdf', '.png', '.jpg', '.jpeg')

_SYSTEM_PROMPT = """당신은 (주)동양 전자결재 검토 전문가입니다. 기안문(본문 + 첨부)을 결재
상신 전에 점검하세요. 실제 문제만 지적하고 "문제 없음/적정" 같은 확인 문장은
findings에 넣지 마세요. 첨부가 있으면 반드시 열어서 내용을 확인하고 판단하세요.
확신이 없으면 지적하지 마세요 — 지적할 때는 반드시 실제로 인용 가능한 근거(quote)를
함께 남기세요(첨부 내용을 인용할 때는 어느 첨부파일인지 알 수 있게 적으세요).
숫자ㆍ금액이 관련된 지적은 특히, 아래 "숫자ㆍ수치 지적 전 검증 원칙"을 반드시 먼저
적용해서 분석한 뒤에만 findings에 넣으세요 — 이 원칙을 건너뛴 지적은 실제로 오탐이었던
사례가 있었습니다.

검토 대상은 [본문]과, 있다면 그 뒤에 오는 [첨부: 파일명] 여러 개입니다.

## 숫자ㆍ수치 지적 전 검증 원칙 (모든 항목에 공통 적용 — 가장 중요)
숫자ㆍ금액이 안 맞는다는 지적은 아래 세 단계를 실제로 거친 뒤에만 findings에 넣으세요.
이 원칙을 어긴 지적(재현 없이 인상만으로 판단, 또는 계산 없이 "이상하다"고 단정)이
과거에 실제로 오탐을 여러 번 냈습니다.

1. **문서가 쓴 기준을 그대로 재현하라.** 문서 문장이 특정 기간ㆍ범위ㆍ대상ㆍ단위ㆍ
   세전후 중 어느 것을 기준으로 판단했는지 먼저 정확히 읽으세요. 검증할 때 그 기준을
   임의로 바꿔치기해서(예: 분기 기준 문장인데 연간 누적으로, 특정 항목 기준인데 전체
   합계로, 세전 금액인데 세후 금액으로) 재계산하면 안 됩니다 — 그건 문서의 오류가
   아니라 검증 자체의 오류입니다. 표에 기간별ㆍ항목별 열이 여러 개 있어도 문서가
   지목한 열만 써야 하고, 문서가 스스로 다른 범위를 기준으로 삼았다고 밝히지 않는 한
   그 범위를 확장해서 비교하지 마세요.
2. **인상이 아니라 실제 계산으로 확인하라.** "달라 보인다/이상하다/단위가 안 맞아
   보인다" 같은 느낌만으로 지적하지 말고, 실제 숫자를 꺼내 덧셈ㆍ나눗셈ㆍ배율 비교
   등을 직접 해본 뒤 그 계산 과정과 결과를 quote/근거에 남기세요. (예: 서로 다른
   단위('원' vs '백만원' 등)로 표시된 여러 시트ㆍ표를 비교할 때는, 대응하는 셀 값을
   골라 선언된 배율만큼 실제로 차이가 나는지 나눗셈으로 검산한 뒤에만 지적하세요.)
3. **계산해서 실제로 앞뒤가 맞으면 지적하지 마라.** 재계산 결과 문서 내용이 맞는데
   범위ㆍ단위ㆍ표현 방식 차이 때문에 처음엔 달라 보였을 뿐이라면, 그 자체가 findings에
   넣지 않아야 할 오탐입니다 — "확인해봤지만 문제없음" 같은 문장도 만들지 말고 그냥
   빼세요.

## 1. 숫자ㆍ금액 정합 (가장 중요)
- 본문과 첨부의 금액ㆍ수량ㆍ날짜ㆍ품번ㆍ업체명이 서로 일치하는지. 첨부끼리도 대조.
- 첨부 안 계산이 맞는지(단가×수량, 소계ㆍ합계, 기초+증가-감소=기말 등).
- 금액마다 부가세 별도/포함이 명시돼 있고, 본문ㆍ계약서ㆍ견적서ㆍ보증금액이 같은
  기준으로 정합한지. (두 금액이 정확히 1.1배 관계면 공급가액 vs VAT포함액이니
  모순이 아닙니다 — 오탐으로 지적하지 마세요.)

## 2. 오탈자 — 실무 오류를 일으키는 것만
- 금액ㆍ숫자ㆍ날짜ㆍ수량 오기, 계약조건 간 모순, 사람ㆍ업체ㆍ거래처명 오기.
- 이전 건(지난달ㆍ지난 분기 등) 문구를 복사해 쓰다가 이번 건과 안 맞게 남은 흔적.
- 띄어쓰기ㆍ맞춤법 등 사소한 표기는 지적하지 마세요.

## 3. 결재자 관점 — 판단 재료가 충분한지
- 무엇을ㆍ왜ㆍ얼마에 하는지, 경위ㆍ리스크ㆍ업체 선정 사유ㆍ승인 요청 범위가 문서에서 읽히는지.
- 판단 근거가 된 서류(견적서ㆍ계약서ㆍ비교표 등)가 첨부에 빠짐없이 있는지.
- 대외로 나가는 공문ㆍ계약서라면, 담아야 할 요청ㆍ조건이 그 문서 자체에 빠지지 않았는지.

## 4. 계약ㆍ업체 서류
- 사업자등록증의 법인명ㆍ사업장이 계약서ㆍ견적서와 일치하는지. 신용조회일이 수개월 지나지 않았는지.
- 보증보험: 보증 종류가 단계에 맞는지(시공 단계=계약이행보증 약 10%, 준공 후=하자보수보증
  약 5%), 가입금액이 VAT 포함 기준으로 부족하지 않은지. (요건보다 크게 가입된 것은 문제 아님)
- 대금지급ㆍ기한ㆍ위약 등 계약 조건이 본문 설명과 어긋나지 않는지.

## 5. 회계ㆍ권한
- 비용의 계정과목ㆍ귀속 사업장이 성격에 맞는지(예: 직원이 아닌 외주기사 관련 비용을
  복리후생비로 계상하면 안 됨).
- 규정상 필요한 직급ㆍ전결권자가 결재선에 빠지지 않았는지.

## 지적하지 않아도 되는 것 (우리 회사 관행상 정상 — 이걸 지적으로 findings에 넣지 마세요)
- 본문은 취지ㆍ요약이고 상세 근거는 첨부에 있는 구조 (본문에 다 안 적혀 있다는 형식 지적 금지)
- 소액 구매나 전속ㆍ독점 품목의 단일견적 / 소액 고정비(시청료ㆍ통신료 등)의 처리 근거 미기재
- 표준 양식의 해당 없는 빈칸 / 발송 전 공문의 문서번호 공란
- 전월 발생 건을 익월 초에 정산ㆍ상신하는 것 (월 마감은 익월 20일경)

## 결과 작성 방법
각 finding은 아래 순서로 판단해서 작성하세요 — ① 결론(무엇이 문제인지) ② 근거
(quote: 원문 위치ㆍ수치) ③ 수정 방법(suggestion). severity는 "critical"(바로 고칠
것 — 명백한 오기ㆍ불일치, 결재 전 반드시 고쳐야 함), "warning"(보완할 것 — 추가할
근거ㆍ첨부ㆍ설명이 필요함, 또는 문서만으로 판단 못해 결재자도 궁금할 확인 질문),
"info"(참고할 만한 사소한 제안) 중 하나로. location에는 "본문", "첨부: 파일명",
"본문-첨부 대조(파일명)" 중 이 지적이 어디에 해당하는지 적으세요.
숫자ㆍ금액 관련 finding은 등록하기 직전에 "숫자ㆍ수치 지적 전 검증 원칙" 3단계(①문서가
쓴 기준 그대로 재현 ②실제 계산으로 확인 ③재계산 결과 맞으면 제외)를 스스로 다시
점검하세요 — 그 결과 실제로는 문제가 없다고 판단되면 findings에서 빼세요.

아래는 검토할 기안 문서입니다."""

_OUTPUT_SCHEMA = {
    "type": "object",
    "properties": {
        "overall_summary": {"type": "string", "description": "전체 검토 결과 1~2문장 요약"},
        "findings": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "location": {"type": "string", "description": '"본문" / "첨부: 파일명" / "본문-첨부 대조(파일명)" 중 하나'},
                    "severity": {"type": "string", "enum": ["critical", "warning", "info"]},
                    "issue": {"type": "string", "description": "구체적인 문제 설명"},
                    "quote": {"type": "string", "description": "근거가 되는 원문 인용"},
                    "suggestion": {"type": "string", "description": "어떻게 수정하면 되는지"},
                },
                "required": ["location", "severity", "issue", "quote", "suggestion"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["overall_summary", "findings"],
    "additionalProperties": False,
}


def _extract_docx_text(file_bytes: bytes) -> str:
    import io
    import docx
    d = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts)


def _extract_xlsx_text(file_bytes: bytes) -> str:
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f'[시트: {ws.title}]')
        for row in ws.iter_rows(values_only=True):
            cells = [str(v) for v in row if v is not None and str(v).strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts)


def _extract_pdf_text(file_bytes: bytes) -> str:
    import fitz  # PyMuPDF
    doc = fitz.open(stream=file_bytes, filetype='pdf')
    parts = [page.get_text() for page in doc]
    doc.close()
    return '\n'.join(parts)


def _safe_filename(name: str) -> str:
    import re
    return re.sub(r'[^\w.\-가-힣]', '_', name)[:150] or 'attachment'


def _save_for_vision(filename: str, file_bytes: bytes, workdir: str, note: str) -> dict:
    import os
    path = os.path.join(workdir, _safe_filename(filename))
    with open(path, 'wb') as f:
        f.write(file_bytes)
    return {'kind': 'file', 'path': path, 'note': note}


IMAGE_EXTS = ('.png', '.jpg', '.jpeg')


def _extract_attachment(filename: str, file_bytes: bytes, workdir: str) -> dict:
    """반환: {'kind': 'text', 'content': str} 또는 {'kind': 'file', 'path': str, 'note': str}
    (텍스트 추출이 안 되는 스캔본 PDF ㆍ 이미지 첨부 — 파일을 workdir에 저장해두고
    Claude가 자기 Read 도구(비전 포함)로 직접 읽게 한다. 본문에 표ㆍ그림이 있어서
    복사-붙여넣기로는 안 옮겨지는 경우, 그 부분만 캡처해서 이미지로 첨부하면 이 경로로
    처리된다 — 그룹웨어 로그인 없이도 해결되는 방법)."""
    ext = ('.' + filename.rsplit('.', 1)[-1].lower()) if '.' in filename else ''
    if ext in IMAGE_EXTS:
        return _save_for_vision(filename, file_bytes, workdir, '이미지 첨부 — Read 도구(비전)로 직접 확인하도록 전달했습니다.')
    try:
        if ext == '.docx':
            text = _extract_docx_text(file_bytes)
        elif ext in ('.xlsx', '.xlsm'):
            text = _extract_xlsx_text(file_bytes)
        elif ext == '.pdf':
            text = _extract_pdf_text(file_bytes).strip()
            if not text:
                return _save_for_vision(
                    filename, file_bytes, workdir,
                    '스캔본(이미지형) PDF로 보여 텍스트 추출이 안 됐습니다 — Read 도구로 직접 읽도록 전달했습니다.',
                )
        else:
            return {'kind': 'text', 'content': f'(지원하지 않는 파일 형식이라 내용을 읽지 못했습니다 — 파일명만 참고: {filename})'}
    except Exception as e:
        return {'kind': 'text', 'content': f'(파일을 읽는 중 오류가 발생해 내용을 넣지 못했습니다: {e})'}

    text = text.strip()
    if not text:
        return {'kind': 'text', 'content': '(추출된 텍스트가 없습니다)'}
    if len(text) > MAX_CHARS_PER_ATTACHMENT:
        text = text[:MAX_CHARS_PER_ATTACHMENT] + f'\n...(이하 생략 — 총 {len(text):,}자 중 앞부분만 포함)'
    return {'kind': 'text', 'content': text}


def _find_claude_cli():
    global _CLAUDE_CLI_PATH
    if _CLAUDE_CLI_PATH == '':
        _CLAUDE_CLI_PATH = shutil.which('claude')
    return _CLAUDE_CLI_PATH


def is_available() -> bool:
    return bool(_find_claude_cli())


def review_text(document_text: str, title: str = '', attachments: list = None, oauth_token: str = None) -> dict:
    """attachments: [(filename, file_bytes), ...] (선택).
    oauth_token: 이 사람 본인이 'claude setup-token'으로 발급받은 개인 토큰(선택) —
    있으면 그 토큰으로(본인 구독 한도), 없으면 이 서버에 로그인된 공용 계정으로 실행한다.
    반환: {"available": bool, "overall_summary": str, "findings": list, "error": str|None,
    "attachment_notes": list}. attachment_notes는 파일별로 텍스트를 잘 읽었는지/생략했는지
    보여주는 참고 정보(오류 여부와 무관하게 항상 채움)."""
    cli = _find_claude_cli()
    if not cli:
        return {
            "available": False, "overall_summary": '', "findings": [], "attachment_notes": [],
            "error": (
                "AI 검수를 쓰려면 이 PC에 Claude Code CLI가 설치·로그인되어 있어야 합니다 — "
                "터미널에서 claude login 을 한 번 실행해주세요. "
                "(방금 설치했다면 통합포털을 재시작해야 인식됩니다.)"
            ),
        }
    if not (document_text or '').strip():
        return {
            "available": True, "overall_summary": '', "findings": [], "attachment_notes": [],
            "error": "검토할 본문 내용이 비어 있습니다 — 기안 문서 내용을 붙여넣어주세요.",
        }

    attachments = attachments or []
    args = [
        cli, '-p', '--output-format', 'json',
        '--json-schema', json.dumps(_OUTPUT_SCHEMA, ensure_ascii=False),
        '--disallowedTools', _CLI_DISALLOWED_TOOLS,
    ]
    with tempfile.TemporaryDirectory(prefix='common_doc_review_') as isolated_cwd:
        # 첨부파일 추출을 subprocess 호출과 같은 임시 폴더 안에서 한다 — 스캔본
        # PDF처럼 텍스트 추출이 안 되는 파일은 이 폴더에 원본을 저장해두고,
        # Claude 자신의 Read 도구(비전 포함)로 직접 읽게 경로를 알려준다.
        attachment_notes = []
        attachment_blocks = []
        for filename, file_bytes in attachments:
            extracted = _extract_attachment(filename, file_bytes, isolated_cwd)
            if extracted['kind'] == 'file':
                attachment_notes.append({'filename': filename, 'preview': extracted['note']})
                attachment_blocks.append(
                    f"[첨부: {filename} — {extracted['note']}]\n"
                    f"Read 도구로 다음 경로의 파일을 직접 읽어 내용을 파악하세요: {extracted['path']}"
                )
            else:
                attachment_notes.append({'filename': filename, 'preview': extracted['content'][:120]})
                attachment_blocks.append(f"[첨부: {filename}]\n{extracted['content']}")

        header = f"문서 제목: {title}\n\n" if title.strip() else ''
        user_prompt = header + "[본문]\n" + document_text
        if attachment_blocks:
            user_prompt += '\n\n' + '\n\n'.join(attachment_blocks)
        combined_prompt = _SYSTEM_PROMPT + '\n\n' + user_prompt

        env = None
        if oauth_token:
            # 이 환경변수가 있으면 claude CLI가 서버에 로그인된 계정 대신 이 토큰
            # (그 사람 개인 구독)으로 인증한다 — 'claude setup-token'으로 발급.
            env = dict(os.environ, CLAUDE_CODE_OAUTH_TOKEN=oauth_token)
        try:
            proc = subprocess.run(
                args, input=combined_prompt, capture_output=True, text=True,
                encoding='utf-8', errors='replace', timeout=CLI_TIMEOUT_SEC, cwd=isolated_cwd, env=env,
            )
        except subprocess.TimeoutExpired:
            return {
                "available": True, "overall_summary": '', "findings": [], "attachment_notes": attachment_notes,
                "error": f'AI 검수가 {CLI_TIMEOUT_SEC}초 안에 끝나지 않아 중단했습니다. 문서가 크거나 응답이 지연되고 있습니다 — 잠시 후 다시 시도해주세요.',
            }
        except FileNotFoundError:
            return {
                "available": True, "overall_summary": '', "findings": [], "attachment_notes": attachment_notes,
                "error": 'Claude Code CLI를 실행하지 못했습니다 — 설치 상태를 다시 확인해주세요.',
            }

    raw = (proc.stdout or '').strip()
    if not raw:
        detail = (proc.stderr or '').strip()
        return {
            "available": True, "overall_summary": '', "findings": [], "attachment_notes": attachment_notes,
            "error": 'AI로부터 응답을 받지 못했습니다.' + (f' ({detail})' if detail else ''),
        }

    try:
        outer = json.loads(raw)
    except json.JSONDecodeError:
        return {
            "available": True, "overall_summary": '', "findings": [], "attachment_notes": attachment_notes,
            "error": 'AI 응답을 해석하지 못했습니다(예상한 JSON 형식이 아닙니다).',
        }

    if outer.get('is_error'):
        msg = str(outer.get('result') or '알 수 없는 오류')
        if 'not logged in' in msg.lower() or '/login' in msg:
            if oauth_token:
                err = (
                    '등록하신 개인 Claude 토큰이 유효하지 않거나 만료된 것 같습니다 — '
                    '내 정보 화면에서 claude setup-token 으로 새로 발급받아 다시 등록해주세요.'
                )
            else:
                err = (
                    'Claude Code CLI에 로그인되어 있지 않습니다 — 통합포털이 돌아가는 이 PC에서 '
                    '터미널을 열고 claude login 을 한 번 실행해 로그인해주세요.'
                )
        else:
            err = f'AI 검수 중 오류: {msg}'
        return {"available": True, "overall_summary": '', "findings": [], "attachment_notes": attachment_notes, "error": err}

    result = outer.get('result')
    if isinstance(result, str):
        try:
            result = json.loads(result)
        except json.JSONDecodeError:
            return {
                "available": True, "overall_summary": '', "findings": [], "attachment_notes": attachment_notes,
                "error": 'AI가 구조화된 형식으로 답하지 않았습니다 — 다시 시도해주세요.',
            }
    if not isinstance(result, dict):
        return {
            "available": True, "overall_summary": '', "findings": [], "attachment_notes": attachment_notes,
            "error": 'AI 응답 형식을 이해하지 못했습니다.',
        }

    return {
        "available": True,
        "overall_summary": result.get("overall_summary", ''),
        "findings": result.get("findings", []),
        "attachment_notes": attachment_notes,
        "error": None,
    }
